"""Render a validated outline into a .pptx.

Claude writes the outline (titles and bullets, each slide citing finding IDs). The renderer checks it,
looks up the source names behind the cited findings, and puts those names and the run date in each
slide's footer. Finding and evidence IDs never appear on a slide.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

SOURCE_NAMES = {
    "dns_records": "Public DNS", "host_resolution": "Public DNS", "certificate_transparency": "Certificate Transparency logs",
    "rdap": "RDAP domain registry", "wayback": "Wayback Machine", "urlscan": "urlscan.io", "m365_tenant": "Microsoft 365 public tenant lookup",
    "cloud_ranges": "Published cloud IP ranges", "ripestat": "RIPEstat", "page_loader": "Public web pages",
    "trust": "Public trust pages", "company": "Public company records", "vulnerabilities": "NVD and CISA KEV",
    "zap_passive": "OWASP ZAP passive scan",
}
ALLOWED_OUTLINE = {"title", "subtitle", "slides"}
ALLOWED_SLIDE = {"title", "bullets", "cites"}
MAX_BULLETS = 7
MAX_SLIDES = 30
# A finding or evidence ID looks like rule:key or source:kind:subject.
ID_LIKE = re.compile(r"\b[a-z][a-z0-9_]*:[a-z0-9_][^\s]*", re.I)


class OutlineError(ValueError):
    pass


def _text(value, where: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise OutlineError(f"{where} must be non-empty text")
    return value.strip()


def validate_outline(outline, bundle: dict) -> dict:
    """Check shape, citations, and that no raw ID leaks into slide text. Returns the outline unchanged."""
    if not isinstance(outline, dict):
        raise OutlineError("outline must be a JSON object")
    extra = set(outline) - ALLOWED_OUTLINE
    if extra:
        raise OutlineError(f"unknown outline field(s): {sorted(extra)}")
    texts = [_text(outline.get("title"), "title")]
    if "subtitle" in outline:
        texts.append(_text(outline["subtitle"], "subtitle"))
    for t in texts:
        if ID_LIKE.search(t):
            raise OutlineError(f"the title page shows a finding or evidence ID: {ID_LIKE.search(t).group(0)}")
    slides = outline.get("slides")
    if not isinstance(slides, list) or not slides:
        raise OutlineError("outline needs at least one slide")
    if len(slides) > MAX_SLIDES:
        raise OutlineError(f"outline has {len(slides)} slides; the limit is {MAX_SLIDES}")
    known = {f["id"] for f in bundle.get("findings", [])}
    for n, slide in enumerate(slides, 1):
        where = f"slide {n}"
        if not isinstance(slide, dict):
            raise OutlineError(f"{where} must be an object")
        extra = set(slide) - ALLOWED_SLIDE
        if extra:
            raise OutlineError(f"{where} has unknown field(s): {sorted(extra)}")
        texts = [_text(slide.get("title"), f"{where} title")]
        bullets = slide.get("bullets", [])
        if not isinstance(bullets, list) or len(bullets) > MAX_BULLETS:
            raise OutlineError(f"{where} bullets must be a list of at most {MAX_BULLETS}")
        texts += [_text(b, f"{where} bullet") for b in bullets]
        cites = slide.get("cites", [])
        if not isinstance(cites, list) or not all(isinstance(c, str) for c in cites):
            raise OutlineError(f"{where} cites must be a list of finding IDs")
        unknown = [c for c in cites if c not in known]
        if unknown:
            raise OutlineError(f"{where} cites findings that are not in the run: {unknown}")
        if bullets and not cites:
            raise OutlineError(f"{where} makes claims but cites no finding")
        for t in texts:
            if ID_LIKE.search(t):
                raise OutlineError(f"{where} shows a finding or evidence ID: {ID_LIKE.search(t).group(0)}")
    return outline


def footer_text(slide: dict, bundle: dict) -> str:
    """Source names behind the cited findings, plus the run date. Never IDs."""
    findings = {f["id"]: f for f in bundle.get("findings", [])}
    source_of = {e["id"]: e["source"] for e in bundle.get("evidence", [])}
    names = []
    for cite in slide.get("cites", []):
        for ev_id in findings[cite]["evidence_ids"]:
            name = SOURCE_NAMES.get(source_of.get(ev_id, ""), "")
            if name and name not in names:
                names.append(name)
    date = str(bundle.get("run_at", ""))[:10]
    return (f"Source: {', '.join(names)}. " if names else "") + f"Run date: {date}"


def _layout(prs, kind: str):
    layouts = list(prs.slide_layouts)
    if kind == "title":
        return layouts[0]
    return layouts[1] if len(layouts) > 1 else layouts[0]


def _add_footer(prs, slide, text: str) -> None:
    width, height = prs.slide_width, prs.slide_height
    box = slide.shapes.add_textbox(Emu(int(width * 0.05)), Emu(int(height * 0.92)), Emu(int(width * 0.9)), Emu(int(height * 0.06)))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for run in frame.paragraphs[0].runs:
        run.font.size = Pt(10)


def render_deck(outline: dict, bundle: dict, out_path: Path, template: Path | None = None) -> Path:
    """Build the .pptx. A template that exists is used for layouts, fonts, and colors."""
    validate_outline(outline, bundle)
    prs = Presentation(str(template)) if template and Path(template).exists() else Presentation()
    if not (template and Path(template).exists()):
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cover = prs.slides.add_slide(_layout(prs, "title"))
    cover.shapes.title.text = outline["title"].strip()
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = str(outline.get("subtitle", "")).strip()
    _add_footer(prs, cover, f"Run date: {str(bundle.get('run_at', ''))[:10]}")
    for item in outline["slides"]:
        slide = prs.slides.add_slide(_layout(prs, "content"))
        slide.shapes.title.text = item["title"].strip()
        body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
        if body is not None:
            frame = body.text_frame
            frame.clear()
            for i, bullet in enumerate(item.get("bullets", [])):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                para.text = bullet.strip()
        _add_footer(prs, slide, footer_text(item, bundle))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def load_json(path: Path) -> dict:
    try:
        return json.loads(Path(path).read_text())
    except (OSError, ValueError) as exc:
        raise OutlineError(f"{path} is not readable JSON: {exc}") from exc
