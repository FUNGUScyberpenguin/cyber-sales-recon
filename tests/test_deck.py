import json

import pytest
from pptx import Presentation

from reconbrief.cli import app
from reconbrief.deck import OutlineError, footer_text, render_deck, validate_outline
from reconbrief.models import Tri
from reconbrief.run import run_recon
from reconbrief.sources.common import make_evidence
from typer.testing import CliRunner

from tests.helpers import FakeHttp, ScriptedDns, response
from tests.rule_helpers import D, host_dns, txt
from tests.test_run_output import HOSTS, OK_HTTP, Stub, ctx_for


@pytest.fixture
def bundle(tmp_path):
    ev = [txt("spf", D, Tri.ABSENT), host_dns("staging.acme.com")]
    run_recon("acme.com", tmp_path, ctx=ctx_for(OK_HTTP), sources=[Stub("dns_records", 1, ev)], preflight_hosts=HOSTS)
    return json.loads((tmp_path / "bundle.json").read_text())


def outline(bundle, **slide):
    fid = next(f["id"] for f in bundle["findings"] if f["rule"] == "non_production_hosts")
    s = {"title": "What we found", "bullets": ["Test systems are public"], "cites": [fid]}
    s.update(slide)
    return {"title": "Acme exposure", "subtitle": "Prepared for the team", "slides": [s]}


def all_text(path):
    prs = Presentation(str(path))
    return [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]


def test_deck_renders_with_source_names_and_run_date_in_footer(bundle, tmp_path):
    path = render_deck(outline(bundle), bundle, tmp_path / "deck.pptx")
    text = " | ".join(all_text(path))
    assert "Test systems are public" in text
    assert "Source: Public DNS. Run date: " + bundle["run_at"][:10] in text


def test_footer_never_shows_ids(bundle, tmp_path):
    path = render_deck(outline(bundle), bundle, tmp_path / "deck.pptx")
    ids = [f["id"] for f in bundle["findings"]] + [e["id"] for e in bundle["evidence"]]
    text = " | ".join(all_text(path))
    assert not any(i in text for i in ids)
    assert "non_production_hosts" not in text


def test_outline_rejects_ids_in_text_unknown_cites_and_extra_fields(bundle):
    fid = bundle["findings"][0]["id"]
    for bad in (outline(bundle, bullets=[f"See {fid}"]), outline(bundle, cites=["made_up:finding"]),
                outline(bundle, cites=[]), outline(bundle, talk_track="say this"), outline(bundle, notes="Tell them X"), outline(bundle, title="")):
        with pytest.raises(OutlineError):
            validate_outline(bad, bundle)
    with pytest.raises(OutlineError):
        validate_outline({**outline(bundle), "discovery_questions": []}, bundle)
    with pytest.raises(OutlineError):
        validate_outline({"title": "x", "slides": []}, bundle)
    with pytest.raises(OutlineError):
        validate_outline({**outline(bundle), "subtitle": f"About {bundle['findings'][0]['id']}"}, bundle)


def test_template_is_used_when_it_exists(bundle, tmp_path):
    template = tmp_path / "brand.pptx"
    base = Presentation()
    base.core_properties.title = "brand"
    base.slide_width = 9144000 // 2
    base.save(str(template))
    path = render_deck(outline(bundle), bundle, tmp_path / "out.pptx", template)
    assert Presentation(str(path)).slide_width == 9144000 // 2
    path = render_deck(outline(bundle), bundle, tmp_path / "out2.pptx", tmp_path / "missing.pptx")
    assert path.exists()


def test_cli_deck_and_run_errors(bundle, tmp_path):
    runner = CliRunner()
    (tmp_path / "bundle.json").write_text(json.dumps(bundle))
    (tmp_path / "outline.json").write_text(json.dumps(outline(bundle)))
    ok = runner.invoke(app, ["deck", str(tmp_path / "outline.json"), "--bundle", str(tmp_path / "bundle.json"), "--out", str(tmp_path / "d.pptx")])
    assert ok.exit_code == 0 and (tmp_path / "d.pptx").exists()
    (tmp_path / "outline.json").write_text(json.dumps(outline(bundle, bullets=["x"], cites=[])))
    bad = runner.invoke(app, ["deck", str(tmp_path / "outline.json"), "--bundle", str(tmp_path / "bundle.json"), "--out", str(tmp_path / "e.pptx")])
    assert bad.exit_code == 2
    assert runner.invoke(app, ["run", "not a domain", "--out", str(tmp_path)]).exit_code == 2


def test_footer_text_without_cites(bundle):
    assert footer_text({}, bundle).startswith("Run date: ")
